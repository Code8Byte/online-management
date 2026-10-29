import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def build_cc_workbook(out_file="LAPORAN_CC_DEPARTEMEN_LIPAT.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan CC"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Page setup for A4 Landscape, fit to 1 page wide and 1 page tall
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Narrow margins (inches)
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2

    # Palettes and Styles
    NAVY_TITLE = "1E3A8A"
    NAVY_HEADER = "1E3A8A"
    ORANGE_TITLE = "C65911"
    ORANGE_HEADER = "ED7D31"
    ORANGE_SUB = "FCE4D6"
    RED_TITLE = "C00000"
    RED_HEADER = "DC2626"
    RED_SUB = "FEE2E2"
    GRAY_TOTAL = "F1F5F9"
    GRAY_BORDER = "CBD5E1"
    TEXT_MAIN = "0F172A"
    TEXT_SUB = "475569"
    RED_CONFIDENTIAL = "DC2626"

    thin_border_side = Side(border_style="thin", color=GRAY_BORDER)
    double_bottom_side = Side(border_style="double", color="0F172A")
    thin_top_side = Side(border_style="thin", color="0F172A")

    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(left=thin_border_side, right=thin_border_side, top=thin_top_side, bottom=double_bottom_side)

    # Fonts (Figtree)
    FONT_FAMILY = "Figtree"
    font_main_title = Font(name=FONT_FAMILY, size=13, bold=True, color=NAVY_TITLE)
    font_sub_title = Font(name=FONT_FAMILY, size=9, bold=False, color=TEXT_SUB)
    font_timestamp = Font(name=FONT_FAMILY, size=8.5, bold=True, color=TEXT_SUB)
    font_section_header = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    font_tbl_header = Font(name=FONT_FAMILY, size=8.5, bold=True, color="FFFFFF")
    font_sub_header = Font(name=FONT_FAMILY, size=7, bold=False, color="7C2D12")
    font_data = Font(name=FONT_FAMILY, size=8.5, bold=False, color=TEXT_MAIN)
    font_data_bold = Font(name=FONT_FAMILY, size=8.5, bold=True, color=TEXT_MAIN)
    font_total = Font(name=FONT_FAMILY, size=8.5, bold=True, color=TEXT_MAIN)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # Fills
    fill_navy_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_orange_title = PatternFill(start_color=ORANGE_TITLE, end_color=ORANGE_TITLE, fill_type="solid")
    fill_orange_header = PatternFill(start_color=ORANGE_HEADER, end_color=ORANGE_HEADER, fill_type="solid")
    fill_orange_sub = PatternFill(start_color=ORANGE_SUB, end_color=ORANGE_SUB, fill_type="solid")
    fill_red_title = PatternFill(start_color=RED_TITLE, end_color=RED_TITLE, fill_type="solid")
    fill_red_header = PatternFill(start_color=RED_HEADER, end_color=RED_HEADER, fill_type="solid")
    fill_total = PatternFill(start_color=GRAY_TOTAL, end_color=GRAY_TOTAL, fill_type="solid")

    now = datetime.now()
    months = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
              "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    export_time_str = f"Dicetak: {now.day} {months[now.month-1]} {now.year}, {now.strftime('%H:%M')} WIB"

    # 1. Page Title & Subtitle & Timestamp (Row 1-2)
    ws.merge_cells("A1:K1")
    ws["A1"] = "LAPORAN KERJA HARIAN - DEPARTEMEN LIPAT"
    ws["A1"].font = font_main_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Timestamp top right
    ws.merge_cells("L1:P1")
    ws["L1"] = export_time_str
    ws["L1"].font = font_timestamp
    ws["L1"].alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = "Tanggal: 19 Oktober 2026   |   Shift: Shift 3   |   Total Mesin: 2 Unit"
    ws["A2"].font = font_sub_title
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 15

    ws.row_dimensions[3].height = 6

    # =========================================================================
    # 2. TABEL 1: TABEL UTAMA (Ringkasan Jam Kerja & Mesin)
    # Tanpa (....) sama sekali!
    # =========================================================================
    t1_headers = [
        ("No", align_center),
        ("Nama Mesin", align_center),
        ("No PP", align_center),
        ("Jam Kerja", align_center),
        ("Total Down Time", align_center),
        ("Jam Jalan", align_center),
        ("Persiapan", align_center)
    ]

    ws.row_dimensions[4].height = 24
    for col_idx, (h_text, h_align) in enumerate(t1_headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h_text)
        cell.font = font_tbl_header
        cell.fill = fill_navy_header
        cell.alignment = h_align
        cell.border = border_data

    # Baris 5: FOCUSIGHT
    ws.row_dimensions[5].height = 20
    ws.cell(row=5, column=1, value=1).alignment = align_center
    ws.cell(row=5, column=2, value="FOCUSIGHT").alignment = align_left
    ws.cell(row=5, column=3, value="7843, 7844").alignment = align_center
    ws.cell(row=5, column=4, value=8.00).alignment = align_right
    ws.cell(row=5, column=4).number_format = "0.00"

    # Formula Total DT links to Tabel 2 Focusight Total DT (N12)
    cell_dt1 = ws.cell(row=5, column=5, value="=N12")
    cell_dt1.alignment = align_right
    cell_dt1.font = font_data_bold
    cell_dt1.number_format = "#,##0"

    # Formula Jam Jalan: =D5 - (E5 + G5)/60
    cell_jj1 = ws.cell(row=5, column=6, value="=D5-(E5+G5)/60")
    cell_jj1.alignment = align_right
    cell_jj1.font = font_data_bold
    cell_jj1.number_format = "0.00"

    # Persiapan Focusight: 30 menit
    cell_p1 = ws.cell(row=5, column=7, value=30)
    cell_p1.alignment = align_right
    cell_p1.number_format = "#,##0"

    for c in range(1, 8):
        ws.cell(row=5, column=c).border = border_data
        if not ws.cell(row=5, column=c).font.bold:
            ws.cell(row=5, column=c).font = font_data

    # Baris 6: LUSTER CARTONHERO-650C
    ws.row_dimensions[6].height = 20
    ws.cell(row=6, column=1, value=2).alignment = align_center
    ws.cell(row=6, column=2, value="LUSTER CARTONHERO-650C").alignment = align_left
    ws.cell(row=6, column=3, value="6059").alignment = align_center
    ws.cell(row=6, column=4, value=8.00).alignment = align_right
    ws.cell(row=6, column=4).number_format = "0.00"

    # Formula Total DT links to Tabel 2 Luster Total DT (N13)
    cell_dt2 = ws.cell(row=6, column=5, value="=N13")
    cell_dt2.alignment = align_right
    cell_dt2.font = font_data_bold
    cell_dt2.number_format = "#,##0"

    # Formula Jam Jalan: =D6 - (E6 + G6)/60
    cell_jj2 = ws.cell(row=6, column=6, value="=D6-(E6+G6)/60")
    cell_jj2.alignment = align_right
    cell_jj2.font = font_data_bold
    cell_jj2.number_format = "0.00"

    # Persiapan Luster: 30 menit
    cell_p2 = ws.cell(row=6, column=7, value=30)
    cell_p2.alignment = align_right
    cell_p2.number_format = "#,##0"

    for c in range(1, 8):
        ws.cell(row=6, column=c).border = border_data
        if not ws.cell(row=6, column=c).font.bold:
            ws.cell(row=6, column=c).font = font_data

    # Baris 7: TOTAL TABEL UTAMA
    ws.row_dimensions[7].height = 20
    ws.merge_cells("A7:C7")
    cell_tot_lbl = ws.cell(row=7, column=1, value="TOTAL")
    cell_tot_lbl.font = font_total
    cell_tot_lbl.alignment = align_center

    ws.cell(row=7, column=4, value="=SUM(D5:D6)").number_format = "0.00"
    ws.cell(row=7, column=5, value="=SUM(E5:E6)").number_format = "#,##0"
    ws.cell(row=7, column=6, value="=SUM(F5:F6)").number_format = "0.00"
    ws.cell(row=7, column=7, value="=SUM(G5:G6)").number_format = "#,##0"

    for c in range(1, 8):
        cell = ws.cell(row=7, column=c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
        if c >= 4:
            cell.alignment = align_right

    ws.row_dimensions[8].height = 8

    # =========================================================================
    # 3. TABEL 2: DETAIL DOWNTIME
    # Tanpa (MENIT)
    # =========================================================================
    ws.row_dimensions[9].height = 20
    ws.merge_cells("A9:N9")
    sec_dt = ws.cell(row=9, column=1, value="DETAIL DOWNTIME")
    sec_dt.font = font_section_header
    sec_dt.fill = fill_orange_title
    sec_dt.alignment = align_center
    for c in range(1, 15):
        ws.cell(row=9, column=c).border = border_data

    # Downtime Codes & Meanings without parentheses
    dt_columns = [
        ("ACCQC", "Acc QC\nModel"),
        ("BDEV", "Ban\nDelivery"),
        ("BEND", "Mesin\nBendel"),
        ("DEFECT", "Perbaikan\nDefect"),
        ("LAIN", "Lain -\nLain"),
        ("PERBK", "Predictive\nRepair"),
        ("SERVC", "Preventive\nService"),
        ("SPC", "Kosongkan\nWIP"),
        ("STEL", "Menyetel\nMesin"),
        ("T.BAHAN", "Tunggu\nBahan")
    ]

    ws.row_dimensions[10].height = 16
    ws.row_dimensions[11].height = 18

    # Base Columns (No, Nama Mesin, No PP) merged vertically Row 10:11
    ws.merge_cells("A10:A11")
    ws.cell(row=10, column=1, value="No").alignment = align_center
    ws.merge_cells("B10:B11")
    ws.cell(row=10, column=2, value="Nama Mesin").alignment = align_center
    ws.merge_cells("C10:C11")
    ws.cell(row=10, column=3, value="No PP").alignment = align_center

    for c in [1, 2, 3]:
        for r in [10, 11]:
            cell = ws.cell(row=r, column=c)
            cell.font = font_tbl_header
            cell.fill = fill_orange_header
            cell.border = border_data

    # DT Columns D to M
    for idx, (kode, arti) in enumerate(dt_columns, start=4):
        c10 = ws.cell(row=10, column=idx, value=kode)
        c10.font = font_tbl_header
        c10.fill = fill_orange_header
        c10.alignment = align_center
        c10.border = border_data
        
        c11 = ws.cell(row=11, column=idx, value=arti)
        c11.font = font_sub_header
        c11.fill = fill_orange_sub
        c11.alignment = align_center
        c11.border = border_data

    # Total DT Column (Col N, Merged N10:N11)
    ws.merge_cells("N10:N11")
    c_tot_dt = ws.cell(row=10, column=14, value="Total DT")
    c_tot_dt.alignment = align_center
    for r in [10, 11]:
        cell = ws.cell(row=r, column=14)
        cell.font = font_tbl_header
        cell.fill = fill_orange_header
        cell.border = border_data

    # Data Downtime Focusight (Row 12)
    dt_focusight = {
        "ACCQC": 0, "BDEV": 0, "BEND": 0, "DEFECT": 30, "LAIN": 0,
        "PERBK": 0, "SERVC": 0, "SPC": 0, "STEL": 60, "T.BAHAN": 60
    }
    ws.row_dimensions[12].height = 20
    ws.cell(row=12, column=1, value=1).alignment = align_center
    ws.cell(row=12, column=2, value="FOCUSIGHT").alignment = align_left
    ws.cell(row=12, column=3, value="7843, 7844").alignment = align_center

    for idx, (kode, _) in enumerate(dt_columns, start=4):
        val = dt_focusight[kode]
        cell = ws.cell(row=12, column=idx, value=val if val > 0 else 0)
        cell.alignment = align_right
        cell.number_format = "#,##0"
        cell.font = font_data_bold if val > 0 else font_data
        cell.border = border_data

    # Formula Total DT Focusight: =SUM(D12:M12)
    c_tot_foc = ws.cell(row=12, column=14, value="=SUM(D12:M12)")
    c_tot_foc.alignment = align_right
    c_tot_foc.font = font_data_bold
    c_tot_foc.number_format = "#,##0"
    c_tot_foc.border = border_data

    for c in range(1, 4):
        ws.cell(row=12, column=c).border = border_data
        ws.cell(row=12, column=c).font = font_data

    # Data Downtime Luster (Row 13)
    dt_luster = {
        "ACCQC": 0, "BDEV": 30, "BEND": 30, "DEFECT": 30, "LAIN": 0,
        "PERBK": 0, "SERVC": 0, "SPC": 0, "STEL": 0, "T.BAHAN": 30
    }
    ws.row_dimensions[13].height = 20
    ws.cell(row=13, column=1, value=2).alignment = align_center
    ws.cell(row=13, column=2, value="LUSTER CARTONHERO-650C").alignment = align_left
    ws.cell(row=13, column=3, value="6059").alignment = align_center

    for idx, (kode, _) in enumerate(dt_columns, start=4):
        val = dt_luster[kode]
        cell = ws.cell(row=13, column=idx, value=val if val > 0 else 0)
        cell.alignment = align_right
        cell.number_format = "#,##0"
        cell.font = font_data_bold if val > 0 else font_data
        cell.border = border_data

    # Formula Total DT Luster: =SUM(D13:M13)
    c_tot_lus = ws.cell(row=13, column=14, value="=SUM(D13:M13)")
    c_tot_lus.alignment = align_right
    c_tot_lus.font = font_data_bold
    c_tot_lus.number_format = "#,##0"
    c_tot_lus.border = border_data

    for c in range(1, 4):
        ws.cell(row=13, column=c).border = border_data
        ws.cell(row=13, column=c).font = font_data

    # Baris 14: TOTAL DETAIL DOWNTIME
    ws.row_dimensions[14].height = 20
    ws.merge_cells("A14:C14")
    c_tot_dt_lbl = ws.cell(row=14, column=1, value="TOTAL")
    c_tot_dt_lbl.font = font_total
    c_tot_dt_lbl.alignment = align_center

    for idx in range(4, 15):
        col_letter = get_column_letter(idx)
        cell = ws.cell(row=14, column=idx, value=f"=SUM({col_letter}12:{col_letter}13)")
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
        cell.alignment = align_right
        cell.number_format = "#,##0"

    for c in range(1, 4):
        cell = ws.cell(row=14, column=c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total

    ws.row_dimensions[15].height = 8

    # =========================================================================
    # 4. TABEL 3: DETAIL WASTE
    # Tanpa (PCS)
    # =========================================================================
    waste_defects = [
        "PREMIER", "SILVER", "WARNA", "UV", "METASOL", "BARET",
        "BLOBOR", "NODA", "EMBOSS", "LAINNYA"
    ]
    total_waste_cols = 3 + 3 + len(waste_defects) # 16 cols
    last_col_letter = get_column_letter(total_waste_cols)

    ws.row_dimensions[16].height = 20
    ws.merge_cells(f"A16:{last_col_letter}16")
    sec_w = ws.cell(row=16, column=1, value="DETAIL WASTE")
    sec_w.font = font_section_header
    sec_w.fill = fill_red_title
    sec_w.alignment = align_center
    for c in range(1, total_waste_cols + 1):
        ws.cell(row=16, column=c).border = border_data

    # Header Tabel 3 (Row 17)
    ws.row_dimensions[17].height = 22
    t3_headers = [
        ("No", align_center),
        ("Nama Mesin", align_center),
        ("No PP", align_center),
        ("Rusakan Stell", align_center),
        ("Reject WIP", align_center),
        ("Total Waste", align_center)
    ]
    for defect_name in waste_defects:
        t3_headers.append((defect_name, align_center))

    for col_idx, (h_text, h_align) in enumerate(t3_headers, start=1):
        cell = ws.cell(row=17, column=col_idx, value=h_text)
        cell.font = font_tbl_header
        cell.fill = fill_red_header
        cell.alignment = h_align
        cell.border = border_data

    # Data Waste Focusight (Row 18)
    w_focusight_defects = {
        "PREMIER": 0, "SILVER": 0, "WARNA": 2750, "UV": 0, "METASOL": 0, "BARET": 0,
        "BLOBOR": 4000, "NODA": 2000, "EMBOSS": 3500, "LAINNYA": 8000
    }
    ws.row_dimensions[18].height = 20
    ws.cell(row=18, column=1, value=1).alignment = align_center
    ws.cell(row=18, column=2, value="FOCUSIGHT").alignment = align_left
    ws.cell(row=18, column=3, value="7843, 7844").alignment = align_center

    ws.cell(row=18, column=4, value=750).number_format = "#,##0" # Stell
    ws.cell(row=18, column=5, value=20250).number_format = "#,##0" # WIP
    c_tot_w1 = ws.cell(row=18, column=6, value="=D18+E18")
    c_tot_w1.number_format = "#,##0"
    c_tot_w1.font = font_data_bold

    for idx, defect_name in enumerate(waste_defects, start=7):
        val = w_focusight_defects[defect_name]
        cell = ws.cell(row=18, column=idx, value=val if val > 0 else 0)
        cell.alignment = align_right
        cell.number_format = "#,##0"
        cell.font = font_data_bold if val > 0 else font_data
        cell.border = border_data

    for c in range(1, 7):
        ws.cell(row=18, column=c).border = border_data
        if c <= 3:
            ws.cell(row=18, column=c).font = font_data
        elif c in [4, 5]:
            ws.cell(row=18, column=c).font = font_data
            ws.cell(row=18, column=c).alignment = align_right

    # Data Waste Luster (Row 19)
    w_luster_defects = {
        "PREMIER": 0, "SILVER": 0, "WARNA": 0, "UV": 2564, "METASOL": 0, "BARET": 0,
        "BLOBOR": 12000, "NODA": 15000, "EMBOSS": 3000, "LAINNYA": 5000
    }
    ws.row_dimensions[19].height = 20
    ws.cell(row=19, column=1, value=2).alignment = align_center
    ws.cell(row=19, column=2, value="LUSTER CARTONHERO-650C").alignment = align_left
    ws.cell(row=19, column=3, value="6059").alignment = align_center

    ws.cell(row=19, column=4, value=850).number_format = "#,##0"
    ws.cell(row=19, column=5, value=37564).number_format = "#,##0"
    c_tot_w2 = ws.cell(row=19, column=6, value="=D19+E19")
    c_tot_w2.number_format = "#,##0"
    c_tot_w2.font = font_data_bold

    for idx, defect_name in enumerate(waste_defects, start=7):
        val = w_luster_defects[defect_name]
        cell = ws.cell(row=19, column=idx, value=val if val > 0 else 0)
        cell.alignment = align_right
        cell.number_format = "#,##0"
        cell.font = font_data_bold if val > 0 else font_data
        cell.border = border_data

    for c in range(1, 7):
        ws.cell(row=19, column=c).border = border_data
        if c <= 3:
            ws.cell(row=19, column=c).font = font_data
        elif c in [4, 5]:
            ws.cell(row=19, column=c).font = font_data
            ws.cell(row=19, column=c).alignment = align_right

    # Baris 20: TOTAL DETAIL WASTE
    ws.row_dimensions[20].height = 20
    ws.merge_cells("A20:C20")
    c_tot_w_lbl = ws.cell(row=20, column=1, value="TOTAL")
    c_tot_w_lbl.font = font_total
    c_tot_w_lbl.alignment = align_center

    for idx in range(4, total_waste_cols + 1):
        col_letter = get_column_letter(idx)
        cell = ws.cell(row=20, column=idx, value=f"=SUM({col_letter}18:{col_letter}19)")
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
        cell.alignment = align_right
        cell.number_format = "#,##0"

    for c in range(1, 4):
        cell = ws.cell(row=20, column=c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total

    ws.row_dimensions[21].height = 12

    # =========================================================================
    # 5. KOLOM TANDA TANGAN (PIC & COST CONTROL)
    # Tanpa tanda kurung!
    # =========================================================================
    ws.row_dimensions[22].height = 18
    ws.merge_cells("B22:D22")
    c_sig_h1 = ws.cell(row=22, column=2, value="Dibuat Oleh,")
    c_sig_h1.font = font_data_bold
    c_sig_h1.alignment = align_center

    ws.merge_cells("G22:J22")
    c_sig_h2 = ws.cell(row=22, column=7, value="Diketahui Oleh,")
    c_sig_h2.font = font_data_bold
    c_sig_h2.alignment = align_center

    # Signature space (Rows 23, 24, 25)
    ws.row_dimensions[23].height = 14
    ws.row_dimensions[24].height = 14
    ws.row_dimensions[25].height = 14

    # Names (Row 26) - Tanpa tanda kurung!
    ws.row_dimensions[26].height = 18
    ws.merge_cells("B26:D26")
    c_sig_n1 = ws.cell(row=26, column=2, value="RENO")
    c_sig_n1.font = Font(name=FONT_FAMILY, size=9, bold=True, underline="single", color=TEXT_MAIN)
    c_sig_n1.alignment = align_center

    ws.merge_cells("G26:J26")
    c_sig_n2 = ws.cell(row=26, column=7, value="____________________________________")
    c_sig_n2.font = Font(name=FONT_FAMILY, size=8.5, bold=False, color=TEXT_SUB)
    c_sig_n2.alignment = align_center

    # Job Titles (Row 27)
    ws.row_dimensions[27].height = 16
    ws.merge_cells("B27:D27")
    c_sig_t1 = ws.cell(row=27, column=2, value="PIC / Operator Sortir")
    c_sig_t1.font = Font(name=FONT_FAMILY, size=8.5, color=TEXT_SUB)
    c_sig_t1.alignment = align_center

    ws.merge_cells("G27:J27")
    c_sig_t2 = ws.cell(row=27, column=7, value="Cost Control")
    c_sig_t2.font = Font(name=FONT_FAMILY, size=8.5, color=TEXT_SUB)
    c_sig_t2.alignment = align_center

    ws.row_dimensions[28].height = 8

    # =========================================================================
    # 6. FOOTER SECTION ON SHEET & IN PRINT SETUP
    # Left: PRO/LKCL-REV_1
    # Center: CONFIDENTIAL (Bold Red)
    # Right: PT PURA BARUTAMA UNIT OFFSET
    # =========================================================================
    ws.row_dimensions[29].height = 18
    line_top = Border(top=Side(style="thin", color="0F172A"))

    ws.merge_cells("B29:D29")
    f_left = ws.cell(row=29, column=2, value="PRO/LKCL-REV_1")
    f_left.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color=TEXT_SUB)
    f_left.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("G29:J29")
    f_cent = ws.cell(row=29, column=7, value="CONFIDENTIAL")
    f_cent.font = Font(name=FONT_FAMILY, size=9, bold=True, color=RED_CONFIDENTIAL)
    f_cent.alignment = align_center

    ws.merge_cells("M29:P29")
    f_right = ws.cell(row=29, column=13, value="PT PURA BARUTAMA UNIT OFFSET")
    f_right.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color=TEXT_SUB)
    f_right.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, total_waste_cols + 1):
        ws.cell(row=29, column=col).border = line_top

    # Setup Excel Print Footer
    ws.oddFooter.left.text = '&"Figtree,Bold"&8PRO/LKCL-REV_1'
    ws.oddFooter.center.text = '&"Figtree,Bold"&9&KDC2626CONFIDENTIAL'
    ws.oddFooter.right.text = '&"Figtree,Bold"&8PT PURA BARUTAMA UNIT OFFSET'

    # Set Print Area
    ws.print_area = f"A1:{last_col_letter}29"

    # =========================================================================
    # 7. COLUMN WIDTHS CONFIGURATION
    # =========================================================================
    column_widths = {
        "A": 5,    # No
        "B": 28,   # Nama Mesin -> Fits "LUSTER CARTONHERO-650C"
        "C": 16,   # No PP -> Fits "7843, 7844"
        "D": 14,   # Jam Kerja / ACCQC / Rusakan Stell
        "E": 18,   # Total Down Time / BDEV / Reject WIP
        "F": 14,   # Jam Jalan / BEND / Total Waste
        "G": 14,   # Persiapan / DEFECT / PREMIER
        "H": 11,   # LAIN / SILVER
        "I": 11,   # PERBK / WARNA
        "J": 11,   # SERVC / UV
        "K": 11,   # SPC / METASOL
        "L": 11,   # STEL / BARET
        "M": 12,   # T.BAHAN / BLOBOR
        "N": 14,   # Total DT / NODA
        "O": 11,   # EMBOSS
        "P": 12    # LAINNYA
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Protect sheet against modification (Read-Only)
    ws.protection.sheet = True
    ws.protection.enable()

    wb.save(out_file)
    print(f"File successfully created: {out_file}")

if __name__ == "__main__":
    build_cc_workbook("LAPORAN_CC_DEPARTEMEN_LIPAT.xlsx")
    build_cc_workbook("LAPORAN_CC_DIVISI_SORTIR_A4_LANDSCAPE.xlsx")
